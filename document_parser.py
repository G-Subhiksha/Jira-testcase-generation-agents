"""
Document Parser
Parses BRD/Story/Epic documents from various formats
"""
import logging
import os
from typing import Dict, Any
from pathlib import Path

class DocumentParser:
    """Parses documents in multiple formats"""
    
    def __init__(self):
        self.logger = logging.getLogger(__name__)
    
    def parse(self, file_path: str) -> Dict[str, Any]:
        """Parse document and return structured content"""
        try:
            file_extension = Path(file_path).suffix.lower()
            
            if file_extension == '.pdf':
                content = self._parse_pdf(file_path)
            elif file_extension in ['.docx', '.doc']:
                content = self._parse_docx(file_path)
            elif file_extension in ['.txt', '.md']:
                content = self._parse_text(file_path)
            else:
                raise ValueError(f"Unsupported file format: {file_extension}")
            
            return {
                "document_content": content,
                "file_path": file_path,
                "file_name": os.path.basename(file_path),
                "file_type": file_extension,
                "character_count": len(content)
            }
            
        except Exception as e:
            self.logger.error(f"Failed to parse document: {str(e)}")
            raise
    
    def _parse_pdf(self, file_path: str) -> str:
        """Parse PDF file"""
        try:
            import pdfplumber
            
            text_content = []
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_content.append(text)
            
            return "\n\n".join(text_content)
            
        except ImportError:
            # Fallback to PyPDF2
            try:
                from PyPDF2 import PdfReader
                
                reader = PdfReader(file_path)
                text_content = []
                
                for page in reader.pages:
                    text = page.extract_text()
                    if text:
                        text_content.append(text)
                
                return "\n\n".join(text_content)
            except ImportError:
                raise ImportError("PDF parsing requires pdfplumber or PyPDF2. Run: pip install pdfplumber")
    
    def _parse_docx(self, file_path: str) -> str:
        """Parse DOCX file"""
        try:
            from docx import Document
            
            doc = Document(file_path)
            text_content = []
            
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_content.append(paragraph.text)
            
            # Extract tables
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    text_content.append(" | ".join(row_text))
            
            return "\n\n".join(text_content)
            
        except ImportError:
            raise ImportError("DOCX parsing requires python-docx. Run: pip install python-docx")
    
    def _parse_text(self, file_path: str) -> str:
        """Parse plain text or markdown file"""
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    
    def parse_from_text(self, content: str, document_type: str = "text") -> Dict[str, Any]:
        """Parse from text content directly"""
        return {
            "document_content": content,
            "file_path": None,
            "file_name": f"pasted_content.{document_type}",
            "file_type": document_type,
            "character_count": len(content)
        }
