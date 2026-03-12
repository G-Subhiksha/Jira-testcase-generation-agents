"""
Excel Export Agent
For EPIC input: Epic Overview | Features | one sheet per Story | Summary | Traceability
For BRD/Story input: Test Cases | Summary | Traceability
"""
from typing import Dict, Any, List
from agents.base_agent import BaseAgent
import json
import os
from datetime import datetime


class ExcelExportAgent(BaseAgent):
    def __init__(self, config: Dict[str, Any], llm_client: Any):
        super().__init__(config, llm_client)
        self.agent_name = "ExcelExportAgent"

    def execute(self, input_data: Dict[str, Any]) -> Dict[str, Any]:
        self.log_start()
        try:
            optimized_suite = input_data.get("optimized_suite", {})
            traceability    = input_data.get("traceability", {})
            classification  = input_data.get("classification", {})
            decomposition   = input_data.get("decomposition", {})   # Epic only
            output_path     = input_data.get("output_path", "output/test_cases.xlsx")

            test_cases = optimized_suite.get("optimizedTestSuite", [])
            self.logger.info(f"Generating Excel workbook with {len(test_cases)} test cases...")

            if decomposition and decomposition.get("features"):
                excel_path = self._export_epic_excel(
                    test_cases, traceability, classification, decomposition, output_path
                )
            else:
                excel_path = self._export_standard_excel(
                    test_cases, traceability, classification, output_path
                )

            self.logger.info(f"Excel file saved to: {excel_path}")
            self.log_end()
            return {
                "excelPath"     : excel_path,
                "testCaseCount" : len(test_cases),
                "agent"         : self.agent_name,
                "timestamp"     : datetime.now().isoformat(),
            }

        except Exception as e:
            self.logger.error(f"Excel export failed: {str(e)}")
            raise

    # =========================================================================
    # EPIC EXCEL  (Epic Overview | Features | per-story sheets | Summary | Traceability)
    # =========================================================================

    def _export_epic_excel(
        self,
        test_cases: List[Dict],
        traceability: Dict,
        classification: Dict,
        decomposition: Dict,
        output_path: str,
    ) -> str:
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("openpyxl not installed. Run: pip3 install openpyxl")

        # ── shared styles ─────────────────────────────────────────────────────
        S = _Styles()

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        features = decomposition.get("features", [])
        epic_title = decomposition.get("epicTitle", "Epic")
        all_stories = [s for f in features for s in f.get("stories", [])]

        # index test cases by storyId for fast lookup
        tc_by_story: Dict[str, List[Dict]] = {}
        for tc in test_cases:
            sid = tc.get("storyId", "ALL")
            tc_by_story.setdefault(sid, []).append(tc)

        # ── Sheet 1: Epic Overview ────────────────────────────────────────────
        ws_epic = wb.create_sheet("Epic Overview")
        self._write_epic_overview(ws_epic, epic_title, classification, features, all_stories, tc_by_story, S)

        # ── Sheet 2: Features ────────────────────────────────────────────────
        ws_feat = wb.create_sheet("Features")
        self._write_features_sheet(ws_feat, features, all_stories, S)

        # ── Sheets 3-N: one per story ─────────────────────────────────────────
        for story in all_stories:
            sid = story.get("storyId", "US-???")
            sheet_name = f"{sid}"   # e.g. "US-001"
            ws_story = wb.create_sheet(sheet_name)
            story_tcs = (
                tc_by_story.get(sid, []) +
                tc_by_story.get("EPIC", [])   # Epic-level security/NF attached to first story
                if story == all_stories[0]
                else tc_by_story.get(sid, [])
            )
            self._write_story_sheet(ws_story, story, story_tcs, S)

        # ── Summary sheet ─────────────────────────────────────────────────────
        ws_sum = wb.create_sheet("Summary")
        self._write_summary_sheet(ws_sum, test_cases, classification, decomposition, S)

        # ── Traceability sheet ────────────────────────────────────────────────
        ws_trace = wb.create_sheet("Traceability Matrix")
        self._write_traceability_sheet(ws_trace, traceability, S)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
        return output_path

    # ── Epic Overview sheet ───────────────────────────────────────────────────

    def _write_epic_overview(self, ws, epic_title, classification, features, stories, tc_by_story, S):
        from openpyxl.styles import Font, Alignment
        ws.column_dimensions["A"].width = 22
        ws.column_dimensions["B"].width = 55
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 14

        r = 1
        # Title banner
        ws.merge_cells(f"A{r}:D{r}")
        c = ws.cell(row=r, column=1, value=f"EPIC: {epic_title}")
        c.fill = S.epic_fill; c.font = Font(bold=True, color="FFFFFF", size=14)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[r].height = 36
        r += 1

        # Metadata rows
        meta = [
            ("Application Type", classification.get("applicationType", "N/A")),
            ("Domain",           classification.get("domain", "N/A")),
            ("Product",          classification.get("product", "N/A")),
            ("Project State",    classification.get("projectState", "N/A")),
            ("Generated On",     datetime.now().strftime("%Y-%m-%d %H:%M")),
            ("Total Features",   len(features)),
            ("Total Stories",    len(stories)),
            ("Total Test Cases", sum(len(v) for v in tc_by_story.values())),
        ]
        for label, value in meta:
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            ws.cell(row=r, column=1).fill = S.alt_fill
            ws.cell(row=r, column=2, value=str(value))
            r += 1

        r += 1
        # Features & stories table
        hdr_vals = ["Feature", "Story", "Priority", "Test Cases"]
        for ci, h in enumerate(hdr_vals, 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.fill = S.header_fill; c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[r].height = 22
        r += 1

        for feat in features:
            first_story_in_feat = True
            for story in feat.get("stories", []):
                sid = story.get("storyId", "")
                tc_count = len(tc_by_story.get(sid, []))
                if first_story_in_feat:
                    ws.cell(row=r, column=1, value=f"{feat['featureId']}: {feat['featureTitle']}").font = Font(bold=True)
                    first_story_in_feat = False
                ws.cell(row=r, column=2, value=f"{sid}: {story.get('title', '')}")
                pri_cell = ws.cell(row=r, column=3, value=story.get("priority", ""))
                pri_cell.fill = S.pri_fill(story.get("priority", ""))
                pri_cell.alignment = Alignment(horizontal="center")
                ws.cell(row=r, column=4, value=tc_count).alignment = Alignment(horizontal="center")
                r += 1

        ws.freeze_panes = "A2"

    # ── Features sheet ────────────────────────────────────────────────────────

    def _write_features_sheet(self, ws, features, all_stories, S):
        from openpyxl.styles import Font, Alignment
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 12
        ws.column_dimensions["D"].width = 30
        ws.column_dimensions["E"].width = 14
        ws.column_dimensions["F"].width = 14

        headers = ["Feature ID", "Feature Title", "Story ID", "Story Title", "Role", "Priority"]
        r = 1
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.fill = S.feat_fill; c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center")
        ws.row_dimensions[r].height = 22
        r += 1

        for feat in features:
            feat_color = PatternFill_import()("solid", fgColor="E8F0FE")
            for si, story in enumerate(feat.get("stories", [])):
                ws.cell(row=r, column=1, value=feat["featureId"] if si == 0 else "").fill = feat_color
                ws.cell(row=r, column=2, value=feat["featureTitle"] if si == 0 else "").fill = feat_color
                ws.cell(row=r, column=3, value=story.get("storyId", ""))
                ws.cell(row=r, column=4, value=story.get("title", ""))
                ws.cell(row=r, column=5, value=story.get("role", ""))
                p = ws.cell(row=r, column=6, value=story.get("priority", ""))
                p.fill = S.pri_fill(story.get("priority", ""))
                p.alignment = Alignment(horizontal="center")
                r += 1

        ws.freeze_panes = "A2"

    # ── Per-story sheet ───────────────────────────────────────────────────────

    def _write_story_sheet(self, ws, story, story_tcs, S):
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter

        thin = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"),  bottom=Side(style="thin"),
        )

        # Column widths
        col_widths = [18, 45, 10, 10, 12, 12, 40, 40, 50, 45, 20]
        col_letters = [get_column_letter(i) for i in range(1, len(col_widths)+1)]
        for ltr, w in zip(col_letters, col_widths):
            ws.column_dimensions[ltr].width = w

        r = 1
        # Story header banner
        ws.merge_cells(f"A{r}:K{r}")
        sid   = story.get("storyId", "")
        title = story.get("title", "")
        c = ws.cell(row=r, column=1, value=f"{sid}: {title}")
        c.fill = S.story_fill; c.font = Font(bold=True, color="FFFFFF", size=12)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[r].height = 30
        r += 1

        # Story metadata
        ws.merge_cells(f"A{r}:K{r}")
        ws.cell(row=r, column=1, value=story.get("story", "")).font = Font(italic=True, size=10)
        ws.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 40
        r += 1

        # Acceptance Criteria
        ws.cell(row=r, column=1, value="Acceptance Criteria").font = Font(bold=True)
        ws.cell(row=r, column=1).fill = S.alt_fill
        r += 1
        for ac in story.get("acceptanceCriteria", []):
            ac_id   = ac.get("acId", "") if isinstance(ac, dict) else ""
            ac_desc = ac.get("description", str(ac)) if isinstance(ac, dict) else str(ac)
            ws.cell(row=r, column=1, value=ac_id).font = Font(bold=True)
            ws.merge_cells(f"B{r}:K{r}")
            ws.cell(row=r, column=2, value=ac_desc).alignment = Alignment(wrap_text=True)
            r += 1

        r += 1  # spacer

        # Test cases header
        tc_headers = ["Test Case ID", "Title", "Category", "Priority",
                      "Automation", "Domain", "Preconditions", "Test Data",
                      "Steps", "Expected Results", "Traceability"]
        for ci, h in enumerate(tc_headers, 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.fill = S.header_fill; c.font = Font(bold=True, color="FFFFFF", size=10)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin
        ws.row_dimensions[r].height = 25
        r += 1

        # Test case rows
        for tc in story_tcs:
            pri = tc.get("priority", "P2")
            cat = tc.get("category", "FUNC")

            preconditions = tc.get("preconditions", [])
            if isinstance(preconditions, list):
                preconditions = "\n".join(f"• {p}" for p in preconditions)

            test_data = tc.get("testData", {})
            if isinstance(test_data, dict):
                test_data = "\n".join(f"{k}: {v}" for k, v in test_data.items())
            elif isinstance(test_data, list):
                test_data = "\n".join(f"• {d}" for d in test_data)

            steps = tc.get("steps", [])
            if isinstance(steps, list):
                steps = "\n".join(f"{i+1}. {s}" for i, s in enumerate(steps))

            expected = tc.get("expectedResults", [])
            if isinstance(expected, list):
                expected = "\n".join(f"✓ {e}" for e in expected)

            row_vals = [
                tc.get("testCaseId", ""), tc.get("title", ""), cat, pri,
                tc.get("automation", ""), tc.get("domain", ""),
                preconditions, test_data, steps, expected,
                tc.get("traceability", ""),
            ]
            cat_fills = {"FUNC": S.func_fill, "API": S.api_fill, "NF": S.nf_fill, "ETL": S.etl_fill}
            row_fill = cat_fills.get(cat, S.alt_fill) if r % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")

            for ci, value in enumerate(row_vals, 1):
                cell = ws.cell(row=r, column=ci, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin
                cell.fill = row_fill
                if ci == 4:   # Priority column
                    cell.fill = S.pri_fill(pri)
                    cell.font = Font(bold=True, color="FFFFFF" if pri in ("P1", "P2") else "000000")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[r].height = max(
                60, 15 * max(1, steps.count("\n") + 1, expected.count("\n") + 1)
            )
            r += 1

        if not story_tcs:
            ws.cell(row=r, column=1, value="No test cases generated for this story.").font = Font(italic=True, color="888888")

        ws.freeze_panes = "A2"

    # ── Summary sheet ─────────────────────────────────────────────────────────

    def _write_summary_sheet(self, ws, test_cases, classification, decomposition, S):
        from openpyxl.styles import Font, Alignment, PatternFill

        ws.column_dimensions["A"].width = 32
        ws.column_dimensions["B"].width = 22

        def hdr(row, text, fill):
            ws.merge_cells(f"A{row}:B{row}")
            c = ws.cell(row=row, column=1, value=text)
            c.fill = fill; c.font = Font(bold=True, color="FFFFFF", size=11)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[row].height = 22

        def row(r, label, value):
            ws.cell(row=r, column=1, value=label)
            ws.cell(row=r, column=2, value=value).alignment = Alignment(horizontal="center")

        cat_counts  = {}
        pri_counts  = {}
        auto_counts = {}
        story_counts = {}
        for tc in test_cases:
            cat_counts [tc.get("category",   "Unknown")] = cat_counts .get(tc.get("category",   "Unknown"), 0) + 1
            pri_counts [tc.get("priority",   "Unknown")] = pri_counts .get(tc.get("priority",   "Unknown"), 0) + 1
            auto_counts[tc.get("automation", "Unknown")] = auto_counts.get(tc.get("automation", "Unknown"), 0) + 1
            sid = tc.get("storyId", "")
            if sid:
                story_counts[sid] = story_counts.get(sid, 0) + 1

        r = 1
        hdr(r, "Test Suite Summary", S.header_fill); r += 1
        row(r, "Total Test Cases",  len(test_cases)); r += 1
        row(r, "Application Type",  classification.get("applicationType", "N/A")); r += 1
        row(r, "Domain",            classification.get("domain", "N/A")); r += 1
        row(r, "Product",           classification.get("product", "N/A")); r += 1
        row(r, "Total Features",    decomposition.get("featureCount", "-")); r += 1
        row(r, "Total Stories",     decomposition.get("storyCount", "-")); r += 1
        row(r, "Generated On",      datetime.now().strftime("%Y-%m-%d %H:%M")); r += 2

        hdr(r, "By Category", S.feat_fill); r += 1
        for k, v in sorted(cat_counts.items()): row(r, k, v); r += 1
        r += 1

        hdr(r, "By Priority", S.story_fill); r += 1
        for k, v in sorted(pri_counts.items()): row(r, k, v); r += 1
        r += 1

        hdr(r, "By Automation Feasibility", S.epic_fill); r += 1
        for k, v in sorted(auto_counts.items()): row(r, k, v); r += 1
        r += 1

        if story_counts:
            hdr(r, "Test Cases per Story", S.header_fill); r += 1
            for sid, cnt in sorted(story_counts.items()):
                row(r, sid, cnt); r += 1

    # ── Traceability sheet ────────────────────────────────────────────────────

    def _write_traceability_sheet(self, ws, traceability, S):
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        from openpyxl.utils import get_column_letter

        thin = Border(left=Side(style="thin"), right=Side(style="thin"),
                      top=Side(style="thin"),  bottom=Side(style="thin"))

        matrix = traceability.get("traceabilityMatrix", [])
        if not matrix:
            ws.cell(row=1, column=1, value="No traceability data.")
            return

        trace_headers = ["Requirement ID", "Story ID", "AC ID", "AC Description", "Test Cases",
                         "Category", "Total Tests", "Coverage Status"]
        trace_widths  = [22, 12, 22, 45, 55, 12, 12, 16]

        r = 1
        for ci, (h, w) in enumerate(zip(trace_headers, trace_widths), 1):
            c = ws.cell(row=r, column=ci, value=h)
            c.fill = S.header_fill; c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.row_dimensions[r].height = 22
        r += 1

        for item in matrix:
            tcs = item.get("testCases", [])
            tcs_str = "\n".join(tcs) if isinstance(tcs, list) else str(tcs)
            row_vals = [
                item.get("requirementId", ""),
                item.get("storyId", ""),
                item.get("acId", ""),
                item.get("acDescription", ""),
                tcs_str,
                ", ".join(item.get("testCasesByCategory", {}).keys()),
                item.get("totalTests", 0),
                item.get("coverage", "Full"),
            ]
            cov = item.get("coverage", "Full")
            if cov == "Full":
                cov_color = PatternFill("solid", fgColor="C6EFCE")
            elif cov == "Partial":
                cov_color = PatternFill("solid", fgColor="FFEB9C")
            else:
                cov_color = PatternFill("solid", fgColor="FFC7CE")

            fill = S.alt_fill if r % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
            for ci, val in enumerate(row_vals, 1):
                cell = ws.cell(row=r, column=ci, value=val)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin
                cell.fill = fill
            # color the coverage cell
            ws.cell(row=r, column=8).fill = cov_color
            ws.row_dimensions[r].height = max(20, 15 * max(1, tcs_str.count("\n") + 1))
            r += 1

        ws.freeze_panes = "A2"

    # =========================================================================
    # STANDARD EXCEL  (BRD / Story — original layout)
    # =========================================================================

    def _export_standard_excel(
        self,
        test_cases: List[Dict],
        traceability: Dict,
        classification: Dict,
        output_path: str,
    ) -> str:
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("openpyxl not installed.")

        S  = _Styles()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Test Cases"

        thin = Border(
            left=Side(style="thin"), right=Side(style="thin"),
            top=Side(style="thin"),  bottom=Side(style="thin"),
        )

        headers    = ["Test Case ID", "Title", "Category", "Priority",
                      "Automation", "Domain", "Project State",
                      "AC Mapped", "Preconditions", "Test Data",
                      "Steps", "Expected Results", "Traceability"]
        col_widths = [18, 45, 10, 10, 12, 20, 15, 12, 40, 40, 50, 45, 20]

        for ci, (h, w) in enumerate(zip(headers, col_widths), 1):
            c = ws.cell(row=1, column=ci, value=h)
            c.fill = S.header_fill; c.font = Font(bold=True, color="FFFFFF", size=11)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = thin
            ws.column_dimensions[get_column_letter(ci)].width = w
        ws.row_dimensions[1].height = 30

        for row_idx, tc in enumerate(test_cases, 2):
            pri = tc.get("priority", "P2")
            cat = tc.get("category", "FUNC")

            preconditions = tc.get("preconditions", [])
            if isinstance(preconditions, list):
                preconditions = "\n".join(f"• {p}" for p in preconditions)
            test_data = tc.get("testData", {})
            if isinstance(test_data, dict):
                test_data = "\n".join(f"{k}: {v}" for k, v in test_data.items())
            elif isinstance(test_data, list):
                test_data = "\n".join(f"• {d}" for d in test_data)
            steps = tc.get("steps", [])
            if isinstance(steps, list):
                steps = "\n".join(f"{i+1}. {s}" for i, s in enumerate(steps))
            expected = tc.get("expectedResults", [])
            if isinstance(expected, list):
                expected = "\n".join(f"✓ {e}" for e in expected)

            values = [
                tc.get("testCaseId", ""), tc.get("title", ""), cat, pri,
                tc.get("automation", ""), tc.get("domain", ""),
                tc.get("projectState", ""), tc.get("acMapped", ""),
                preconditions, test_data, steps, expected,
                tc.get("traceability", ""),
            ]
            cat_fills = {"FUNC": S.func_fill, "API": S.api_fill, "NF": S.nf_fill, "ETL": S.etl_fill}
            row_fill = cat_fills.get(cat, S.alt_fill) if row_idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")

            for ci, value in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=ci, value=value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = thin
                cell.fill = row_fill
                if ci == 4:
                    cell.fill = S.pri_fill(pri)
                    cell.font = Font(bold=True, color="FFFFFF" if pri in ("P1", "P2") else "000000")
                    cell.alignment = Alignment(horizontal="center", vertical="center")

            ws.row_dimensions[row_idx].height = max(
                60, 15 * max(1, steps.count("\n") + 1, expected.count("\n") + 1)
            )

        ws.freeze_panes = "A2"

        # Summary sheet
        ws2 = wb.create_sheet("Summary")
        self._write_summary_sheet(ws2, test_cases, classification, {}, S)

        # Traceability sheet
        ws3 = wb.create_sheet("Traceability Matrix")
        self._write_traceability_sheet(ws3, traceability, S)

        os.makedirs(os.path.dirname(output_path), exist_ok=True)
        wb.save(output_path)
        return output_path


# =============================================================================
# Shared style helpers
# =============================================================================

def PatternFill_import():
    from openpyxl.styles import PatternFill
    return PatternFill


class _Styles:
    def __init__(self):
        from openpyxl.styles import PatternFill, Font, Alignment
        self.header_fill = PatternFill("solid", fgColor="1F4E79")
        self.epic_fill   = PatternFill("solid", fgColor="203864")
        self.feat_fill   = PatternFill("solid", fgColor="2E75B6")
        self.story_fill  = PatternFill("solid", fgColor="2F5597")
        self.alt_fill    = PatternFill("solid", fgColor="F2F2F2")
        self.func_fill   = PatternFill("solid", fgColor="DDEEFF")
        self.api_fill    = PatternFill("solid", fgColor="E6F3E6")
        self.nf_fill     = PatternFill("solid", fgColor="FFF0E6")
        self.etl_fill    = PatternFill("solid", fgColor="F0E6FF")
        self._p1 = PatternFill("solid", fgColor="C00000")
        self._p2 = PatternFill("solid", fgColor="E26B0A")
        self._p3 = PatternFill("solid", fgColor="FFD700")
        self._p4 = PatternFill("solid", fgColor="70AD47")
        self._ph = PatternFill("solid", fgColor="C00000")
        self._pm = PatternFill("solid", fgColor="E26B0A")
        self._pl = PatternFill("solid", fgColor="70AD47")

    def pri_fill(self, pri: str):
        m = {
            "P1": self._p1, "P2": self._p2, "P3": self._p3, "P4": self._p4,
            "High": self._ph, "Medium": self._pm, "Low": self._pl,
        }
        return m.get(pri, self._p3)
